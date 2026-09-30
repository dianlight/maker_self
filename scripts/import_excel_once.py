"""Import una tantum da Inventario.xlsx -> CSV per Sheet + preview JSON."""
import argparse
import csv
import json
import re

CANDIDATES = {
    "codice": ["codice"],
    "quantita": ["quantità", "quantita"],
    "descrizione": ["decription", "description", "descrizione"],
    "interfaccia": ["interface", "interfaccia"],
    "note": ["note"],
    "datasheet": ["pinout/datasheet", "pinout", "datasheet"],
    "posizione": ["position", "posizione"],
}


def slug(s):
    s = (s or "").lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-") or "na"


def norm_header(h):
    return (h or "").strip().lower()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--out", default="import_preview.json")
    ap.add_argument("--csv", default="import_for_sheet.csv")
    a = ap.parse_args()
    from openpyxl import load_workbook
    wb = load_workbook(a.xlsx, data_only=True, read_only=True)
    rows_out = []
    for ws in wb.worksheets:
        if ws.title.lower().startswith("export"):
            continue
        header = None
        for row in ws.iter_rows(values_only=True):
            vals = [(v or "") for v in row]
            joined = " ".join(str(v).lower() for v in vals)
            if "quantit" in joined and ("decription" in joined or "descri" in joined):
                header = [norm_header(v) for v in row]
                break
        if not header:
            continue
        pos = {}
        for key, names in CANDIDATES.items():
            for i, h in enumerate(header):
                if h in names:
                    pos[key] = i
                    break
        for row in ws.iter_rows(values_only=True):
            if row == tuple(header) or row[0] is None and all(v is None for v in row):
                continue
            vals = list(row) + [None] * (len(header) - len(row))
            get = lambda k: str(vals[pos[k]] or "").strip() if k in pos else ""
            try:
                q = int(float(get("quantita") or 0))
            except ValueError:
                q = 0
            if not (get("codice") or get("descrizione")):
                continue
            cat = ws.title.strip().lower().replace(" ", "-").replace("/", "-")
            rows_out.append({
                "id": "", "categoria": cat, "codice": get("codice"),
                "quantita": max(q, 0), "descrizione": get("descrizione"),
                "interfaccia": get("interfaccia"), "note": get("note"),
                "datasheet_url": get("datasheet"), "posizione": get("posizione"),
                "foto_drive_id": "", "foto_url": "", "ai_proposta": "",
                "ai_stato": "", "updated_at": "",
            })
    with open(a.out, "w") as f:
        json.dump(rows_out, f, ensure_ascii=False, indent=2)
    cols = ["id", "categoria", "codice", "quantita", "descrizione", "interfaccia",
            "note", "datasheet_url", "posizione", "foto_drive_id", "foto_url",
            "ai_proposta", "ai_stato", "updated_at"]
    with open(a.csv, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(rows_out)
    print(f"{len(rows_out)} righe -> {a.out}, {a.csv}")


if __name__ == "__main__":
    main()
